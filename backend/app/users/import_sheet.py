from __future__ import annotations

import csv
import io
from dataclasses import dataclass

from app.common.exceptions import AppError


NAME_HEADERS = {"name", "full_name", "full name", "employee name", "employee"}
EMAIL_HEADERS = {"email", "email id", "email_id", "e-mail", "mail"}
PASSWORD_HEADERS = {"password", "pass", "pwd"}
ROLE_HEADERS = {"role", "roles"}


@dataclass(frozen=True, slots=True)
class ImportedUserRow:
    row_number: int
    full_name: str
    email: str
    password: str
    role: str | None = None


def parse_user_sheet(filename: str, content: bytes) -> list[ImportedUserRow]:
    name = (filename or "").lower()
    if name.endswith(".csv") or name.endswith(".txt"):
        return _parse_csv(content)
    if name.endswith(".xlsx"):
        return _parse_xlsx(content)
    raise AppError(
        "UNSUPPORTED_FILE",
        "Upload an Excel workbook (.xlsx) or CSV file with columns name, email, password.",
        400,
    )


def _normalize_header(value: object) -> str:
    return " ".join(str(value or "").strip().lower().replace("_", " ").split())


def _pick(row: dict[str, str], keys: set[str]) -> str:
    for key, value in row.items():
        if key in keys:
            return value.strip()
    return ""


def _rows_from_dicts(dicts: list[dict[str, str]], start_row: int = 2) -> list[ImportedUserRow]:
    parsed: list[ImportedUserRow] = []
    for index, raw in enumerate(dicts, start=start_row):
        name = _pick(raw, NAME_HEADERS)
        email = _pick(raw, EMAIL_HEADERS).lower()
        password = _pick(raw, PASSWORD_HEADERS)
        role = _pick(raw, ROLE_HEADERS) or None
        if not name and not email and not password:
            continue
        parsed.append(
            ImportedUserRow(
                row_number=index,
                full_name=name,
                email=email,
                password=password,
                role=role,
            )
        )
    if not parsed:
        raise AppError("EMPTY_SHEET", "No user rows were found in the file.", 400)
    return parsed


def _parse_csv(content: bytes) -> list[ImportedUserRow]:
    text = content.decode("utf-8-sig", errors="replace")
    reader = csv.DictReader(io.StringIO(text))
    if not reader.fieldnames:
        raise AppError("INVALID_SHEET", "The CSV file is missing a header row.", 400)
    rows: list[dict[str, str]] = []
    for item in reader:
        rows.append({_normalize_header(key): str(value or "") for key, value in item.items()})
    return _rows_from_dicts(rows)


def _parse_xlsx(content: bytes) -> list[ImportedUserRow]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise AppError(
            "UNSUPPORTED_FILE",
            "Excel import needs openpyxl. Save the sheet as CSV, or install openpyxl on the API.",
            400,
        ) from exc
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        sheet = workbook.active
        if sheet is None:
            raise AppError("INVALID_SHEET", "The Excel file has no worksheet.", 400)
        iterator = sheet.iter_rows(values_only=True)
        try:
            headers = [_normalize_header(cell) for cell in next(iterator)]
        except StopIteration as exc:
            raise AppError("INVALID_SHEET", "The Excel file is empty.", 400) from exc
        dicts: list[dict[str, str]] = []
        for values in iterator:
            dicts.append(
                {
                    header: "" if value is None else str(value).strip()
                    for header, value in zip(headers, values, strict=False)
                    if header
                }
            )
        return _rows_from_dicts(dicts)
    finally:
        workbook.close()
